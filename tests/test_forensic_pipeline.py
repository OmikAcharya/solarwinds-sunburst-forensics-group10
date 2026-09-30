"""
test_forensic_pipeline.py
Automated Pytest Suite for Digital Forensics Triage Pipeline & Statutory Compliance.

Department of Computer Engineering, KJ Somaiya School of Engineering
Somaiya Vidyavihar University, Mumbai.
Digital Forensics & Cyber Security Laboratory (Capstone Activity - Group 10)
Authors:
  - Amandeep Singh (Roll No: 16010123036) — Lead Investigator
  - Omik Acharya   (Roll No: 16010123218) — Reverse Engineer
  - Om Lanke       (Roll No: 16010123216) — Cyber Legal Auditor
"""

import os
import json
import pytest
import hashlib
from pathlib import Path

# Add project root to sys.path so scripts can be imported
import sys
PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from scripts.run_forensics import (
    SUNBURST_SHA256_PRIMARY,
    SUNBURST_SHA256_SECONDARY,
    SUNBURST_MD5,
    SUNBURST_SHA1,
    SUNBURST_SIZE,
    SUNBURST_NAME,
    run_pipeline,
    compute_hashes,
    decompress_sunburst_string,
    fnv1a_64,
    check_process_blacklist,
    simulate_environment_behavior,
    ENCODED_STRINGS_DATASET
)
from scripts.generate_bsa_cert import (
    generate_certificate_text,
    get_system_telemetry,
    DEFAULT_EVIDENCE
)


class TestCryptographicIntegrity:
    """Verifies baseline cryptographic hashes and hash verification mechanisms."""

    def test_sample_hash_manifest_presence_and_format(self):
        hash_file = PROJECT_ROOT / "evidence" / "sample_hash.sha256"
        assert hash_file.exists(), "evidence/sample_hash.sha256 manifest must exist"
        
        content = hash_file.read_text(encoding="utf-8").strip()
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        assert len(lines) >= 1, "Hash manifest must contain at least one hash entry"
        
        hashes_found = [line.split()[0] for line in lines]
        assert SUNBURST_SHA256_PRIMARY in hashes_found

    def test_evidence_metadata_schema_and_values(self):
        meta_file = PROJECT_ROOT / "evidence" / "evidence_metadata.json"
        assert meta_file.exists(), "evidence/evidence_metadata.json must exist"
        
        data = json.loads(meta_file.read_text(encoding="utf-8"))
        assert "case_metadata" in data
        assert "artifact_metadata" in data

        artifact = data["artifact_metadata"]
        assert artifact["file_size_bytes"] == SUNBURST_SIZE
        assert artifact["hashes"]["sha256"] == SUNBURST_SHA256_PRIMARY
        assert artifact["hashes"]["md5"] == SUNBURST_MD5
        assert artifact["hashes"]["sha1"] == SUNBURST_SHA1
        assert "SolarWinds Worldwide, LLC" in artifact["authenticode_signature"]["signer_identity"]
        assert artifact["threat_intelligence"]["primary_c2_dga_apex"] == "avsvmcloud.com"


class TestDeobfuscationAndEvasionMechanics:
    """Validates Deflate-raw string decoding and FNV-1a 64-bit + XOR process hashing."""

    def test_deflate_string_decompression(self):
        # Test C2 apex extraction
        c2_plain = decompress_sunburst_string("SywrLstNzskvTdFLzs8FAA==")
        assert c2_plain == "avsvmcloud.com"

        # Test Registry fingerprint path
        reg_plain = decompress_sunburst_string("C/Z3Cwl3DHKN8c1MLsovzk8riXEuqiwoyU8vSizIqAQA")
        assert "SOFTWARE\\Microsoft\\Cryptography" in reg_plain

        # Test MachineGuid value
        guid_plain = decompress_sunburst_string("801MzsjMS3UvzUwBAA==")
        assert guid_plain == "MachineGuid"

    def test_fnv1a_64_hashing_and_xor_blocklist(self):
        # Wireshark test
        res_ws = check_process_blacklist("Wireshark")
        assert res_ws["is_blacklisted"] is True
        assert res_ws["fnv1a_hex"] == "0xa84ff6500970f54d"
        assert res_ws["xor_val"] == 17574002783607647274

        # Non-blacklisted benign tool
        res_notepad = check_process_blacklist("notepad")
        assert res_notepad["is_blacklisted"] is False

    def test_behavior_flow_simulation_environments(self):
        # Victim environment should pass all checks
        victim_sim = simulate_environment_behavior("victim")
        assert "persistent C2 channel" in victim_sim["final_verdict"]

        # Analyst environment should terminate at Step 4
        analyst_sim = simulate_environment_behavior("analyst")
        assert "Stopped at Step 4" in analyst_sim["final_verdict"]


class TestForensicArtifactParsers:
    """Validates the structure and indicators in DiE, FLOSS, and CAPA reports."""

    def test_die_inspection_report_contents(self):
        die_file = PROJECT_ROOT / "outputs" / "die_inspection.txt"
        assert die_file.exists(), "outputs/die_inspection.txt must exist"
        
        content = die_file.read_text(encoding="utf-8")
        assert "Detect It Easy" in content
        assert "PE32" in content
        assert ".NET CLR METADATA HEADER" in content
        assert "v4.0.30319" in content
        assert "SolarWinds Worldwide, LLC" in content
        assert "Entropy:            6.21" in content or "6.21" in content

    def test_floss_decoded_strings_content(self):
        floss_file = PROJECT_ROOT / "outputs" / "floss_decoded_strings.txt"
        assert floss_file.exists(), "outputs/floss_decoded_strings.txt must exist"
        
        content = floss_file.read_text(encoding="utf-8")
        assert "FLOSS (FLARE Obfuscated String Solver)" in content
        assert "sysmon.exe" in content
        assert "wireshark.exe" in content
        assert "processhacker.exe" in content
        assert "x64dbg.exe" in content
        assert "avsvmcloud.com" in content

    def test_capa_capabilities_report_att_and_ck_mapping(self):
        capa_file = PROJECT_ROOT / "outputs" / "capa_capabilities_report.txt"
        assert capa_file.exists(), "outputs/capa_capabilities_report.txt must exist"
        
        content = capa_file.read_text(encoding="utf-8")
        assert "capa" in content
        assert "T1497.003" in content  # Time-based delay
        assert "T1562.001" in content  # Impair defenses
        assert "T1027" in content      # Obfuscated strings
        assert "T1071.004" in content  # DNS C2


class TestStatutoryAndLegalCompliance:
    """Verifies generation and validity of legal compliance certificates."""

    def test_system_telemetry_gathering(self):
        telemetry = get_system_telemetry()
        assert "hostname" in telemetry
        assert "mac_address" in telemetry
        assert len(telemetry["mac_address"].split(":")) == 6

    def test_bsa_certificate_content_and_affirmations(self):
        cert_md = generate_certificate_text()
        assert "SECTION 63 OF THE BHARATIYA SAKSHYA ADHINIYAM (BSA), 2023" in cert_md
        assert "Section 63(4)(c)" in cert_md
        assert "Bharatiya Nagarik Suraksha Sanhita" in cert_md
        assert "Amandeep Singh" in cert_md
        assert "16010123036" in cert_md
        assert "Omik Acharya" in cert_md
        assert "16010123218" in cert_md
        assert "Om Lanke" in cert_md
        assert "16010123216" in cert_md
        assert SUNBURST_SHA256_PRIMARY in cert_md

    def test_cert_in_template_compliance(self):
        cert_in_file = PROJECT_ROOT / "legal_compliance" / "CERT_In_Incident_Notification_Template.md"
        assert cert_in_file.exists(), "CERT_In incident template must exist"
        content = cert_in_file.read_text(encoding="utf-8")
        assert "CERT-In MANDATORY CYBER SECURITY INCIDENT REPORTING FORM" in content
        assert "Section 70B(6)" in content
        assert "180 days" in content or "180-Day" in content or "180-day" in content
        assert "avsvmcloud.com" in content


class TestPipelineExecution:
    """Tests runtime simulation and standalone execution of forensic engine."""

    def test_run_forensics_simulation_mode(self, tmp_path):
        out_tmp = tmp_path / "test_outputs"
        result = run_pipeline(target_path=None, outdir=str(out_tmp), simulate=True)
        assert result["status"] == "success"
        assert result["hash_match"] is True
        assert (out_tmp / "die_inspection.txt").exists()
        assert (out_tmp / "floss_decoded_strings.txt").exists()
        assert (out_tmp / "capa_capabilities_report.txt").exists()
        assert (out_tmp / "fnv_blocklist_hashes.txt").exists()
        assert (out_tmp / "behavior_flow_reconstruction.txt").exists()


class TestDocumentationDeliverables:
    """Asserts that all required documentation files exist and are substantial."""

    @pytest.mark.parametrize("rel_path,min_bytes", [
        ("docs/Investigation_Report.md", 5000),
        ("docs/Presentation_Script.md", 5000),
        ("docs/Presentation_Slides_Outline.md", 3000),
        ("docs/sunburst_simulation.html", 10000),
        ("legal_compliance/Section_63_BSA_Certificate.md", 3000),
        ("legal_compliance/CERT_In_Incident_Notification_Template.md", 2000),
        ("evidence/chain_of_custody.md", 1500),
        ("scripts/verify_hashes.sh", 500),
        ("scripts/generate_bsa_cert.py", 1000),
        ("scripts/run_forensics.py", 2000),
        ("scripts/build_demo_docs.py", 2000),
        ("scripts/build_presentation.py", 2000),
        ("docs/Tool_Execution_and_Demo_Manual.md", 3000),
        ("docs/Presentation_Speaking_Script.docx", 20000),
        ("docs/Tool_Execution_and_Demo_Manual.docx", 20000),
        ("docs/SUNBURST_Forensic_Investigation_Group10.pptx", 30000),
    ])
    def test_deliverable_file_integrity(self, rel_path, min_bytes):
        file_path = PROJECT_ROOT / rel_path
        assert file_path.exists(), f"Required deliverable missing: {rel_path}"
        file_size = file_path.stat().st_size
        assert file_size >= min_bytes, f"Deliverable {rel_path} too small ({file_size} < {min_bytes} bytes)"

    def test_presentation_deck_integrity(self):
        import pptx
        deck_path = PROJECT_ROOT / "docs" / "SUNBURST_Forensic_Investigation_Group10.pptx"
        assert deck_path.exists(), "SUNBURST PPTX deck must exist"
        prs = pptx.Presentation(str(deck_path))
        assert len(prs.slides) == 12, f"Presentation must contain exactly 12 slides, found {len(prs.slides)}"
